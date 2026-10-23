from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from auto_eval.web import server
from auto_eval.web.comparison_uploads import ComparisonUploadStore
from auto_eval.web.operation_comparison_import import import_operation_comparison_file


def _content(correctness: str = "ok") -> bytes:
    return json.dumps({
        "id": "001", "query": "打开设置", "correctness": correctness,
        "issue_types": [], "rationale": "画面已打开设置", "标签": "原始标签",
        "metadata": {"id": "原始来源ID"}, "video_path": "data/example.mp4",
    }, ensure_ascii=False).encode("utf-8")


@pytest.fixture
def upload_client(tmp_path, monkeypatch):
    store = ComparisonUploadStore(tmp_path / "uploads")
    monkeypatch.setattr(server, "comparison_upload_store", store)
    with TestClient(server.app) as client:
        yield client, store


def _upload(client, name: str = "原始文件.jsonl", *, compact: bool = True):
    response = client.post(
        f"/api/operation/comparison/import?compact={str(compact).lower()}",
        files={"file": (name, _content(), "application/json")},
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_store_survives_recreation_and_preserves_full_source(tmp_path: Path):
    root = tmp_path / "uploads"
    source = import_operation_comparison_file("原文件.jsonl", _content())
    source["warnings"] = [{"row": i, "message": f"警告{i}"} for i in range(25)]
    source["summary"]["warning_count"] = 25
    summary = ComparisonUploadStore(root).save(source)
    assert "rows" not in summary
    assert len(summary["warnings"]) == 20
    reloaded = ComparisonUploadStore(root).load(source["source_id"])
    assert reloaded["rows"] == source["rows"]
    assert reloaded["warnings"] == source["warnings"]
    assert reloaded["rows"][0]["source_data"]["标签"] == "原始标签"
    assert reloaded["rows"][0]["source_data"]["metadata"]["id"] == "原始来源ID"


def test_history_lists_only_summaries_and_clamps_deleted_last_page(tmp_path: Path):
    store = ComparisonUploadStore(tmp_path / "uploads")
    for i in range(12):
        source = import_operation_comparison_file(f"组{i}.jsonl", _content())
        summary = store.save(source)
        # 若列表误读逐题文件，这里会抛异常。
        (store.root / summary["source_id"] / "source.json").write_text("broken", encoding="utf-8")
    first = store.list_page()
    second = store.list_page(2)
    assert first["total"] == 12
    assert len(first["items"]) == 10
    assert len(second["items"]) == 2
    assert first["items"][0]["created_at"] >= second["items"][0]["created_at"]
    for item in second["items"]:
        store.delete(item["source_id"])
    assert store.list_page(2)["page"] == 1


def test_failed_save_is_not_visible_and_existing_upload_is_not_overwritten(tmp_path: Path):
    store = ComparisonUploadStore(tmp_path / "uploads")
    source = import_operation_comparison_file("原文件.jsonl", _content())
    source["rows"][0]["source_data"]["bad"] = float("nan")
    with pytest.raises(ValueError):
        store.save(source)
    assert store.list_page()["total"] == 0
    assert list(store.root.iterdir()) == []
    del source["rows"][0]["source_data"]["bad"]
    store.save(source)
    with pytest.raises(ValueError, match="不允许覆盖"):
        store.save(source)
    assert store.list_page()["total"] == 1


def test_upload_history_api_persists_edits_and_keeps_raw_data(upload_client):
    client, store = upload_client
    source = _upload(client)
    assert source["stored"] is True
    assert "rows" not in source
    source_id = source["source_id"]
    base = f"/api/operation/comparison/uploads/{source_id}"
    assert client.get("/api/operation/comparison/uploads").json()["total"] == 1
    assert client.patch(base, json={"note": "第一行\n第二行"}).status_code == 200
    updated = client.patch(base, json={"group_name": "留档名称"}).json()
    assert updated["note"] == "第一行\n第二行"
    assert updated["filename"] == "原始文件.jsonl"
    assert client.get(base).json()["group_name"] == "留档名称"
    persisted = ComparisonUploadStore(store.root).load(source_id)
    assert persisted["group_name"] == "留档名称"
    assert persisted["rows"][0]["source_data"]["标签"] == "原始标签"
    assert not list(store.root.glob("**/.writing_*"))
    assert client.delete(base).status_code == 200
    assert client.get(base).status_code == 404
    assert client.delete(base).status_code == 404
    assert client.get("/api/operation/comparison/uploads").json()["total"] == 0


def test_default_upload_response_keeps_legacy_rows(upload_client):
    client, _ = upload_client
    source = _upload(client, compact=False)
    assert source["rows"][0]["id"] == "001"
    assert source["stored"] is True
    assert client.get("/api/operation/comparison/uploads").json()["total"] == 1


def test_repeated_upload_creates_independent_history_and_failed_import_is_not_saved(upload_client):
    client, _ = upload_client
    first, second = _upload(client), _upload(client)
    assert first["source_id"] != second["source_id"]
    response = client.post(
        "/api/operation/comparison/import",
        files={"file": ("无效.csv", b"query\nexample", "text/csv")},
    )
    assert response.status_code == 422
    assert client.get("/api/operation/comparison/uploads").json()["total"] == 2


def test_saved_references_support_analysis_and_both_exports_after_restart(upload_client, monkeypatch):
    client, store = upload_client
    first = _upload(client, "对照.jsonl")
    second = _upload(client, "实验.jsonl")
    client.patch(f"/api/operation/comparison/uploads/{first['source_id']}", json={"group_name": "已保存对照"})
    monkeypatch.setattr(server, "comparison_upload_store", ComparisonUploadStore(store.root))
    request = {
        "sources": [{"source_id": row["source_id"], "source_type": "upload", "stored": True}
                    for row in (first, second)],
        "control_source_id": first["source_id"], "scope": "common_decidable",
    }
    analyzed = client.post("/api/operation/comparison/analyze", json=request)
    assert analyzed.status_code == 200, analyzed.text
    payload = analyzed.json()
    assert payload["groups"][0]["group_name"] == "已保存对照"
    assert all(group["source_type"] == "upload" for group in payload["groups"])
    exported = client.post("/api/operation/comparison/export?format=xlsx", json=request)
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content))
    assert workbook.sheetnames[:2] == ["逐题横向对比", "数据集明细"]
    values = list(workbook["数据集明细"].values)
    assert any("原始标签" in row for row in values)
    assert client.post("/api/operation/comparison/export?format=html", json=request).status_code == 200
    client.delete(f"/api/operation/comparison/uploads/{first['source_id']}")
    # 不回退到客户端提供的旧行；删除引用应明确报错。
    request["sources"][0]["rows"] = import_operation_comparison_file("旧.jsonl", _content())["rows"]
    assert client.post("/api/operation/comparison/analyze", json=request).status_code == 404


def test_invalid_ids_pagination_and_oversized_notes_are_rejected(upload_client):
    client, _ = upload_client
    assert client.get("/api/operation/comparison/uploads/bad-id").status_code == 422
    assert client.delete("/api/operation/comparison/uploads/bad-id").status_code == 422
    assert client.get("/api/operation/comparison/uploads?page=0").status_code == 422
    assert client.get("/api/operation/comparison/uploads?page_size=1000").status_code == 422
    source = _upload(client)
    assert client.patch(
        f"/api/operation/comparison/uploads/{source['source_id']}", json={"note": "长" * 10001},
    ).status_code == 422
