from __future__ import annotations

import json
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from auto_eval.analysis.operation_comparison import compare_operation_batches
from auto_eval.analysis.operation_report import build_comparison_report
from auto_eval.web.history import build_operation_comparison_xlsx, operation_comparison_batch
from auto_eval.web.operation_comparison_import import (
    import_operation_comparison_file,
    validate_uploaded_comparison_source,
)


def _snapshot(task_id: str, records: list[dict]) -> dict:
    return {
        "task_id": task_id, "dataset_name": f"{task_id}.jsonl", "mode": "operation",
        "status": "done", "options": {},
        "items": [
            {"id": record["id"], "query": record["query"], "source_data": dict(record)}
            for record in records
        ],
        "results": [
            {"index": index, "correctness": "ok", "rationale": "完成"}
            for index in range(len(records))
        ],
    }


def _batches(*snapshots: dict) -> list[dict]:
    return [operation_comparison_batch(snapshot, include_source_data=True) for snapshot in snapshots]


def _sheet_rows(sheet) -> list[dict]:
    headers = [cell.value for cell in sheet[2]]
    return [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=3, values_only=True)]


def test_shared_dataset_is_one_row_per_case_in_horizontal_order() -> None:
    control = _snapshot("control", [
        {"id": "b", "query": "关闭定位", "context": "实验室", "业务标签": "设备设置",
         "metadata": {"id": "business-b", "priority": 0},
         "对比序号": "原始序号", "对比序号_1": "原始标签"},
        {"id": "a", "query": "打开蓝牙", "custom_value": False},
    ])
    experiment = _snapshot("experiment", [
        {"id": "a", "query": "打开蓝牙"},
        {"id": "b", "query": "关闭定位", "context": "实验室", "业务标签": "另一组标签"},
        {"id": "c", "query": "打开相机", "extra_label": "实验组独有题"},
    ])
    payload = compare_operation_batches(
        _batches(control, experiment), baseline_task_id="control", include_union_rows=True,
    )
    workbook = load_workbook(BytesIO(build_operation_comparison_xlsx(payload)), data_only=True)
    assert workbook.sheetnames[:2] == ["逐题横向对比", "数据集明细"]
    results, dataset = workbook.worksheets[:2]
    assert results.max_row == dataset.max_row == 5
    assert [results.cell(row, 1).value for row in range(3, 6)] == [1, 2, 3]
    rows = _sheet_rows(dataset)
    assert [row["对比序号"] for row in rows] == [1, 2, 3]
    assert [row["id"] for row in rows] == ["b", "a", "c"]
    assert [row["明细来源组"] for row in rows] == ["对照组", "对照组", "实验组A"]
    assert rows[0]["业务标签"] == "设备设置"
    assert rows[0]["metadata.id"] == "business-b"
    assert json.loads(rows[0]["metadata"]) == {"id": "business-b", "priority": 0}
    assert rows[0]["对比序号_2"] == "原始序号"
    assert rows[0]["对比序号_1"] == "原始标签"
    assert rows[1]["custom_value"] == "False"
    assert rows[2]["extra_label"] == "实验组独有题"
    # 第二个表第 5 行对应第一表第 5 行：对照组整段为空，实验组记录仍在。
    assert results.cell(5, 5).value in (None, "")
    assert payload["all_groups_common_valid_count"] == 2


@pytest.mark.parametrize("duplicate", [False, True])
def test_conflicting_or_duplicate_ids_keep_separate_dataset_rows(duplicate: bool) -> None:
    records = [{"id": "same", "query": "打开蓝牙"}]
    if duplicate:
        records.append({"id": "same", "query": "关闭定位"})
    control = _snapshot("control", records)
    experiment = _snapshot("experiment", [{"id": "same", "query": "关闭定位"}])
    payload = compare_operation_batches(
        _batches(control, experiment), baseline_task_id="control", include_union_rows=True,
    )
    workbook = load_workbook(BytesIO(build_operation_comparison_xlsx(payload)))
    rows = _sheet_rows(workbook["数据集明细"])
    assert len(rows) == len(records) + 1
    assert len({row["对比序号"] for row in rows}) == len(rows)
    assert all(row["匹配键"] == "same" for row in rows)
    assert workbook["逐题横向对比"].max_row == len(rows) + 2
    assert payload["all_groups_common_valid_count"] == 0


def test_shared_input_differences_are_visible_without_changing_statistics() -> None:
    control = _snapshot("control", [{
        "id": "a", "query": "识别图片", "context": "地点甲", "attachment_path": "/input/a.png",
        "video_path": "/control.mp4", "private_label": "SOURCE_ONLY",
    }])
    experiment = _snapshot("experiment", [{
        "id": "a", "query": "识别图片", "context": "地点乙", "query_images": ["/input/b.png"],
        "video_path": "/experiment.mp4",
    }])
    batches = _batches(control, experiment)
    payload = compare_operation_batches(batches, baseline_task_id="control", include_union_rows=True)
    assert payload["all_groups_common_valid_count"] == 1
    assert payload["alignment_summary"]["input_mismatch"] == 1
    assert payload["alignment_warnings"][0]["fields"] == ["context", "query_images"]
    workbook = load_workbook(BytesIO(build_operation_comparison_xlsx(payload)))
    row = _sheet_rows(workbook["数据集明细"])[0]
    assert row["context"] == "地点甲"
    assert row["共享输入差异"] == "实验组A：context、query_images"
    assert "SOURCE_ONLY" not in json.dumps(build_comparison_report(batches, payload))
    compact = [operation_comparison_batch(snapshot) for snapshot in (control, experiment)]
    assert all("source_data" not in batch["rows"][0] for batch in compact)
    assert compare_operation_batches(compact, baseline_task_id="control")["alignment_summary"]["input_mismatch"] == 1


def test_shared_dataset_keeps_failed_and_pending_cases_outside_statistical_scope() -> None:
    records = [
        {"id": "a", "query": "打开蓝牙", "原始标签": "失败题"},
        {"id": "b", "query": "关闭定位", "原始标签": "待评估题"},
    ]
    control = _snapshot("control", records)
    control["results"] = [{"index": 0, "error": "模型超时"}]
    experiment = _snapshot("experiment", records)
    payload = compare_operation_batches(
        _batches(control, experiment), baseline_task_id="control",
        include_union_rows=True, scope="common_decidable",
    )
    assert payload["all_groups_common_valid_count"] == 0
    workbook = load_workbook(BytesIO(build_operation_comparison_xlsx(payload)))
    assert workbook["逐题横向对比"].max_row == workbook["数据集明细"].max_row == 4
    assert [row["原始标签"] for row in _sheet_rows(workbook["数据集明细"])] == ["失败题", "待评估题"]


def test_optional_inputs_and_image_aliases_do_not_cause_false_differences() -> None:
    control = _snapshot("control", [{
        "id": "a", "query": "识别图片", "attachment_path": "/input/a.png", "context": "实验室",
    }])
    experiment = _snapshot("experiment", [{
        "id": "a", "query": "识别图片", "query_images": '["/input/a.png"]',
    }])
    payload = compare_operation_batches(_batches(control, experiment), baseline_task_id="control")
    assert payload["alignment_summary"]["input_mismatch"] == 0
    assert payload["alignment_warnings"] == []


@pytest.mark.parametrize("format", ["csv", "xlsx", "jsonl"])
def test_uploaded_custom_fields_survive_validation_and_shared_export(format: str) -> None:
    record = {
        "id": "001", "query": "打开蓝牙", "correctness": "ok", "issue_types": "",
        "rationale": "已完成", "任意业务标签": "实验数据", "context": "实验室",
    }
    if format == "jsonl":
        record["metadata"] = {"id": "business-001"}
        content = json.dumps(record).encode()
    elif format == "csv":
        content = pd.DataFrame([record]).to_csv(index=False).encode()
    else:
        output = BytesIO()
        pd.DataFrame([record]).to_excel(output, sheet_name="评估结果", index=False)
        content = output.getvalue()
    uploaded = import_operation_comparison_file(f"input.{format}", content)
    batch = validate_uploaded_comparison_source(uploaded)
    assert batch["rows"][0]["source_data"]["任意业务标签"] == "实验数据"
    assert batch["rows"][0]["source_data"]["id"] == "001"
    another = {**batch, "task_id": "other"}
    payload = compare_operation_batches([batch, another], baseline_task_id=batch["task_id"], include_union_rows=True)
    workbook = load_workbook(BytesIO(build_operation_comparison_xlsx(payload)))
    row = _sheet_rows(workbook["数据集明细"])[0]
    assert row["任意业务标签"] == "实验数据"
    if format == "jsonl":
        assert row["metadata.id"] == "business-001"


def test_full_excel_recovers_dataset_columns_by_identity_not_row_order() -> None:
    result_records = [
        {"id": item_id, "query": query, "correctness": "ok", "issue_types": "", "rationale": "已完成"}
        for item_id, query in [("a", "打开蓝牙"), ("b", "关闭定位"), ("c", "打开设置")]
    ]
    output = BytesIO()
    with pd.ExcelWriter(output) as writer:
        pd.DataFrame(result_records).to_excel(writer, sheet_name="评估结果", index=False)
        pd.DataFrame([
            {"id": "b", "query": "关闭定位", "原始标签": "B"},
            {"id": "a", "query": "打开蓝牙", "原始标签": "A"},
            {"id": "c", "query": "不同问题", "原始标签": "不能误关联"},
        ]).to_excel(writer, sheet_name="数据集明细", index=False)
    uploaded = import_operation_comparison_file("full.xlsx", output.getvalue())
    assert uploaded["rows"][0]["source_data"]["原始标签"] == "A"
    assert uploaded["rows"][1]["source_data"]["原始标签"] == "B"
    assert "原始标签" not in uploaded["rows"][2]["source_data"]
    assert uploaded["summary"]["valid_count"] == 3
    assert uploaded["warnings"][0]["row"] == 4
