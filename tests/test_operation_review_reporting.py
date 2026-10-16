"""效果问题和需复核原因的导入、配对统计、报告及导出使用同一数据链路。"""
from __future__ import annotations

import copy
import json
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from auto_eval.analysis.operation_comparison import compare_operation_batches
from auto_eval.analysis.operation_report import build_comparison_report
from auto_eval.web.history import build_operation_comparison_xlsx
from auto_eval.web.operation_comparison_import import import_operation_comparison_file


@pytest.mark.parametrize("format,nested", [("jsonl", False), ("jsonl", True), ("csv", False), ("xlsx", False)])
def test_import_preserves_review_categories(format: str, nested: bool) -> None:
    evaluation = {
        "correctness": "no_support", "issue_types": "内部过程信息泄露",
        "need_review_cate": "待用户澄清；待用户澄清；缺少前置条件",
    }
    record = {"id": "a", "query": "查询信息"}
    if nested:
        record["evaluation"] = evaluation
    else:
        record.update(evaluation)
        record["需复核类别"] = record.pop("need_review_cate")
    if format == "jsonl":
        content = json.dumps(record, ensure_ascii=False).encode()
    elif format == "csv":
        content = pd.DataFrame([record]).to_csv(index=False).encode()
    else:
        output = BytesIO()
        pd.DataFrame([record]).to_excel(output, index=False)
        content = output.getvalue()
    source = import_operation_comparison_file(f"review.{format}", content)
    result = source["rows"][0]["result"]
    assert result["need_review_cate"] == ["待用户澄清", "缺少前置条件"]
    assert result["issue_types"] == ["内部过程信息泄露"]
    assert source["rows"][0]["export"]["need_review_cate"] == result["need_review_cate"]
    assert source["mapping"]["need_review_cate"] == ("evaluation.need_review_cate" if nested else "需复核类别")


def test_import_legacy_others_moves_review_without_losing_effect_issue() -> None:
    source = import_operation_comparison_file("legacy.jsonl", json.dumps({
        "id": "a", "query": "查询信息", "correctness": "others",
        "issue_types": ["环境或算力异常", "内部过程信息泄露"],
    }, ensure_ascii=False).encode())
    result = source["rows"][0]["result"]
    assert source["summary"]["valid_count"] == 1
    assert result["correctness"] == "no_support"
    assert result["judge_correctness"] == "others"
    assert result["need_review_cate"] == ["环境或算力异常"]
    assert result["issue_types"] == ["内部过程信息泄露"]


def review_batches() -> list[dict]:
    judgments = [
        [
            dict(correctness="ok", issue_types=["回复语义重复"], need_review_cate=["关键事实待核验"]),
            dict(correctness="no_support", issue_types=["回复语义重复"], need_review_cate=["待用户澄清"] * 2),
            dict(correctness="others", issue_types=["环境或算力异常"]),
            dict(correctness="nok", issue_types=["任务结果错误"], need_review_cate=["缺少前置条件"]),
            dict(correctness="ok"),
        ],
        [
            dict(correctness="nok", issue_types=["任务结果错误"]),
            dict(correctness="ok", issue_types=["回复语义重复"]),
            dict(correctness="no_support", need_review_cate=["环境或算力异常"]),
            dict(correctness="no_support", need_review_cate=["待用户澄清"]),
            dict(correctness="ok", error="timeout"),
        ],
    ]
    return [
        dict(task_id=name, dataset_name=name, rows=[
            dict(id=str(index), item_id=str(index), query=f"q{index}", result=result)
            for index, result in enumerate(results)
        ]) for name, results in zip(("a", "b"), judgments)
    ]


@pytest.mark.parametrize("scope", ["all_valid", "common_decidable"])
def test_review_pairing_uses_all_valid_even_when_effect_scope_filters(scope: str) -> None:
    batches = review_batches()
    original = copy.deepcopy(batches)
    payload = compare_operation_batches(batches, baseline_task_id="a", scope=scope)
    pair = payload["pairwise"][0]
    assert pair["valid_pair_count"] == (4 if scope == "all_valid" else 1)
    assert pair["review_valid_pair_count"] == 4
    assert pair["baseline_review_case_count"] == 4
    assert pair["target_review_case_count"] == 2
    rows = {row["category"]: row for row in pair["review_category_rows"]}
    assert rows["待用户澄清"]["baseline_count"] == 1
    assert rows["待用户澄清"]["target_rate"] == .25
    assert rows["缺少前置条件"]["rate_delta"] == -.25
    assert rows["环境或算力异常"]["count_delta"] == 0
    assert all(row["issue_type"] not in rows for row in pair["issue_type_rows"])
    if scope == "all_valid":
        repeated = next(row for row in pair["issue_type_rows"] if row["issue_type"] == "回复语义重复")
        assert repeated["baseline_count"] == 1  # no_support 的质量问题保留在明细，但不计效果分布。
        assert repeated["baseline_rate"] == .5
    report = build_comparison_report(batches, payload)
    assert len(report["pairs"][0]["review_matches"]) == 4
    assert len(report["pairs"][0]["matches"]) == pair["valid_pair_count"]
    assert report["pairs"][0]["review_category_rows"] == pair["review_category_rows"]
    assert report["groups"][0]["cases"][2]["correctness"] == "no_support"
    assert report["groups"][0]["cases"][2]["issue_types"] == []
    assert report["groups"][0]["cases"][2]["need_review_cate"] == ["环境或算力异常"]
    assert batches == original


def test_xlsx_review_sheet_is_separate_and_neutral() -> None:
    payload = compare_operation_batches(review_batches(), baseline_task_id="a", include_union_rows=True, scope="common_decidable")
    workbook = load_workbook(BytesIO(build_operation_comparison_xlsx(payload)))
    sheet = workbook["需复核类别对比"]
    rows = [dict(zip([cell.value for cell in sheet[1]], values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
    assert all(row["共同有效Case"] == 4 for row in rows)
    assert next(row for row in rows if row["需复核类别"] == "缺少前置条件")["占比差值（pp）"] == -25
    assert all(sheet.cell(row=index, column=8).fill.patternType is None for index in range(2, sheet.max_row + 1))
    assert sheet.freeze_panes == "A2"
    assert sheet["E2"].number_format == "0.00%"
    detail = workbook["逐题横向对比"]
    assert [cell.value for cell in detail[2]].count("need_review_cate") == 2
    assert "环境或算力异常" in str(list(detail.values))
    assert "对照组效果分母" in [cell.value for cell in workbook["Issue Type对比"][1]]


def test_empty_review_pair_has_no_fake_percentage() -> None:
    batches = review_batches()
    for batch in batches:
        for row in batch["rows"]:
            row["result"] = {"error": "failed"}
    pair = compare_operation_batches(batches, baseline_task_id="a")["pairwise"][0]
    assert pair["review_valid_pair_count"] == 0
    assert pair["review_category_rows"] == []
    assert pair["baseline_review_case_count"] == pair["target_review_case_count"] == 0
